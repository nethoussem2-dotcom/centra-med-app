import openpyxl
import os
import sqlite3
import datetime
from database import init_db, get_connection, add_user
from auth import hash_password

# Paths
EXCEL_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), "SUIVI_21.05.2026.xlsx")
DB_PATH = os.path.join(os.path.dirname(__file__), "centra_med.db")

def seed_users():
    """Seeds default admin, operator, and visitor accounts."""
    print("Seeding users...")
    conn = get_connection()
    cursor = conn.cursor()
    
    # Clean users table first to avoid duplication
    cursor.execute("DELETE FROM users")
    conn.commit()
    conn.close()
    
    # Add users
    # Admin
    admin_pw = hash_password("admin123")
    add_user("admin", admin_pw, "مدير النظام CENTRA MED", "admin")
    
    # Operator / User
    user_pw = hash_password("user123")
    add_user("operator", user_pw, "مستلم الإنتاج (أمين)", "user")
    
    # Visitor
    visitor_pw = hash_password("visitor123")
    add_user("visitor", visitor_pw, "زائر عام (اطلاع فقط)", "visitor")
    
    print("Users seeded successfully:")
    print("  - Admin: admin / admin123")
    print("  - User: operator / user123")
    print("  - Visitor: visitor / visitor123")

def seed_products(wb):
    """Parses 'ART PF' sheet and seeds the art_pf table."""
    print("\nSeeding products from 'ART PF' sheet...")
    if "ART PF" not in wb.sheetnames:
        print("Error: 'ART PF' sheet not found in Excel file.")
        return
        
    sheet = wb["ART PF"]
    rows = list(sheet.iter_rows(values_only=True))
    
    # Find header row index
    header_row_index = -1
    for idx, r in enumerate(rows[:10]):
        if r and ('CODE ART PF' in r or 'CODE ARTICLE' in r):
            header_row_index = idx
            break
            
    if header_row_index == -1:
        # Hardcode fallback to Row 0
        header_row_index = 0
        
    conn = get_connection()
    cursor = conn.cursor()
    
    # Clear products table
    cursor.execute("DELETE FROM art_pf")
    
    headers = [str(h).upper().strip() if h is not None else "" for h in rows[header_row_index]]
    
    try:
        code_idx = next(i for i, h in enumerate(headers) if 'CODE ART PF' in h or 'CODE ARTICLE' in h)
        desig_idx = next(i for i, h in enumerate(headers) if 'DESIGNATION' in h)
        poids_idx = next(i for i, h in enumerate(headers) if 'POID/PCE' in h or 'POIDS' in h)
        std_dechet_idx = next(i for i, h in enumerate(headers) if 'DECHET' in h or 'STANDAR' in h)
        cadence_idx = next(i for i, h in enumerate(headers) if 'CADANCE' in h or 'CADENCE' in h)
        matiere_idx = next(i for i, h in enumerate(headers) if 'D MP' in h or 'MATIERE' in h)
        um_idx = next(i for i, h in enumerate(headers) if h == 'UM')
    except StopIteration:
        # Precise indices fallback
        code_idx = 1
        desig_idx = 2
        poids_idx = 3
        std_dechet_idx = 4
        cadence_idx = 5
        matiere_idx = 6
        um_idx = 22
        
    print(f"Product headers mapped: Code={code_idx}, Desig={desig_idx}, Weight={poids_idx}, Waste={std_dechet_idx}, Cadence={cadence_idx}, Material={matiere_idx}, UM={um_idx}")
    
    product_count = 0
    for r in rows[header_row_index + 1:]:
        if not r or len(r) <= max(code_idx, desig_idx):
            continue
            
        code = r[code_idx]
        designation = r[desig_idx]
        
        # Skip if code or designation is empty
        if not code or not designation or code == "CODE ART PF" or str(code).strip() == "":
            continue
            
        # Extract fields
        code = str(code).strip()
        designation = str(designation).strip()
        
        try:
            poids = float(r[poids_idx]) if r[poids_idx] is not None else 0.0
        except ValueError:
            poids = 0.0
            
        try:
            std_dechet = float(r[std_dechet_idx]) if r[std_dechet_idx] is not None else 0.0
        except ValueError:
            std_dechet = 0.0
            
        try:
            cadence = float(r[cadence_idx]) if r[cadence_idx] is not None else None
        except ValueError:
            cadence = None
                
        matiere = str(r[matiere_idx]).strip() if r[matiere_idx] is not None else ""
        um = str(r[um_idx]).strip() if r[um_idx] is not None else "PCE"
        
        cursor.execute("""
        INSERT OR REPLACE INTO art_pf (code_article, designation, poids_theorique_gr, standard_dechet, cadence_theo, matiere_premiere, um)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (code, designation, poids, std_dechet, cadence, matiere, um))
        product_count += 1
        
    conn.commit()
    conn.close()
    print(f"Products seeded successfully: {product_count} products added.")

def seed_production_records(wb):
    """Parses 'SUIVI PROD2026 JORNAL' and seeds the production_records table."""
    print("\nSeeding production records...")
    
    sheet_name = "SUIVI PROD2026 JORNAL"
    if sheet_name not in wb.sheetnames:
        sheet_name = "SUIVI J PROD 2024"
        if sheet_name not in wb.sheetnames:
            print("Error: Neither 2026 nor 2024 production sheets found.")
            return
            
    sheet = wb[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    
    # Find header row
    header_row_index = -1
    for idx, r in enumerate(rows[:15]):
        if r and any(str(cell).upper().strip() == 'CODE ARTICLE' for cell in r if cell is not None):
            header_row_index = idx
            break
            
    if header_row_index == -1:
        header_row_index = 4
        
    headers = [str(h).upper().strip() if h is not None else "" for h in rows[header_row_index]]
    
    # Map index defaults or dynamically search
    try:
        date_idx = next(i for i, h in enumerate(headers) if 'DATE' in h)
        code_idx = next(i for i, h in enumerate(headers) if 'CODE ARTICLE' in h)
        machine_idx = next(i for i, h in enumerate(headers) if 'MACHINE' in h)
        desig_idx = next(i for i, h in enumerate(headers) if 'DESIGNATION' in h)
        lot_idx = next(i for i, h in enumerate(headers) if 'LOT' in h)
        qte_idx = next(i for i, h in enumerate(headers) if 'QTE' in h or 'PCS' in h)
        poid_theo_idx = next(i for i, h in enumerate(headers) if 'POID THEO' in h or 'THEORIQUE' in h)
        poid_reel_idx = next(i for i, h in enumerate(headers) if 'POID REEL' in h or 'REEL' in h)
        dechet_pce_idx = next(i for i, h in enumerate(headers) if 'DECHET (PCE)' in h or 'DECHET(PCS)' in h or ('DECHET' in h and 'PC' in h))
        dechet_kg_idx = next(i for i, h in enumerate(headers) if 'DECHET (KG)' in h or 'DECHET(KG)' in h or ('DECHET' in h and 'KG' in h))
        rebut_idx = next(i for i, h in enumerate(headers) if 'REBUT' in h or 'TAUX' in h)
    except StopIteration:
        # Default indexes
        date_idx = 0
        code_idx = 1
        machine_idx = 2
        desig_idx = 3
        lot_idx = 4
        qte_idx = 7
        poid_theo_idx = 8
        poid_reel_idx = 9
        dechet_pce_idx = 10
        dechet_kg_idx = 11
        rebut_idx = 12
        
    print(f"Production headers mapped: Date={date_idx}, Code={code_idx}, Machine={machine_idx}, Desig={desig_idx}, Lot={lot_idx}, Qte={qte_idx}, PoidReel={poid_reel_idx}, DechetKg={dechet_kg_idx}")
    
    conn = get_connection()
    cursor = conn.cursor()
    
    # Clear production_records table
    cursor.execute("DELETE FROM production_records")
    
    record_count = 0
    skipped_count = 0
    
    for r in rows[header_row_index + 1:]:
        if not r or len(r) <= max(date_idx, code_idx, desig_idx):
            continue
            
        date_val = r[date_idx]
        code_art = r[code_idx]
        machine = r[machine_idx]
        designation = r[desig_idx]
        num_lot = r[lot_idx]
        
        # Skip empty lines or header rows
        if not date_val or not code_art or str(code_art).strip() == "" or str(date_val).upper().strip() == "DATE":
            skipped_count += 1
            continue
            
        # Parse Date
        if isinstance(date_val, datetime.datetime):
            date_str = date_val.strftime("%Y-%m-%d")
        elif isinstance(date_val, datetime.date):
            date_str = date_val.strftime("%Y-%m-%d")
        else:
            try:
                date_str = str(date_val).split(" ")[0].strip()
                datetime.datetime.strptime(date_str, "%Y-%m-%d")
            except ValueError:
                date_str = datetime.date.today().strftime("%Y-%m-%d")
                
        code_art = str(code_art).strip()
        machine = str(machine).strip() if machine else "M01"
        designation = str(designation).strip()
        num_lot = str(num_lot).strip() if num_lot else "N/A"
        
        # Parse Numeric Values
        try:
            qte_prod = float(r[qte_idx]) if r[qte_idx] is not None else 0.0
        except ValueError:
            qte_prod = 0.0
            
        try:
            poid_theo = float(r[poid_theo_idx]) if r[poid_theo_idx] is not None else 0.0
        except ValueError:
            poid_theo = 0.0
            
        try:
            poid_reel = float(r[poid_reel_idx]) if r[poid_reel_idx] is not None else 0.0
        except ValueError:
            poid_reel = 0.0
            
        try:
            dechet_pce = float(r[dechet_pce_idx]) if r[dechet_pce_idx] is not None else 0.0
        except ValueError:
            dechet_pce = 0.0
            
        try:
            dechet_kg = float(r[dechet_kg_idx]) if r[dechet_kg_idx] is not None else 0.0
        except ValueError:
            dechet_kg = 0.0
            
        try:
            taux_rebut = float(r[rebut_idx]) if r[rebut_idx] is not None else 0.0
        except ValueError:
            taux_rebut = 0.0
            
        # Re-calculate values if zero or default to ensure database consistency
        cursor.execute("SELECT poids_theorique_gr FROM art_pf WHERE code_article = ?", (code_art,))
        p_row = cursor.fetchone()
        if p_row:
            p_weight = p_row[0]
            if p_weight > 0:
                if poid_theo == 0.0 and qte_prod > 0:
                    poid_theo = qte_prod * (p_weight / 1000.0)
                if dechet_pce == 0.0 and dechet_kg > 0:
                    dechet_pce = dechet_kg / (p_weight / 1000.0)
                    
        if taux_rebut == 0.0 and dechet_kg > 0:
            total_w = poid_reel + dechet_kg
            if total_w > 0:
                taux_rebut = dechet_kg / total_w
                    
        # Assign to default operator
        username = "operator"
        
        cursor.execute("""
        INSERT INTO production_records (
            date, username, code_article, designation, machine, num_lot, qte_prod, poid_theo, poid_reel, dechet_pce, dechet_kg, taux_rebut
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (date_str, username, code_art, designation, machine, num_lot, qte_prod, poid_theo, poid_reel, dechet_pce, dechet_kg, taux_rebut))
        record_count += 1
        
    conn.commit()
    conn.close()
    print(f"Production records seeded successfully: {record_count} logs added. Skipped lines: {skipped_count}")

def main():
    print("=== CENTRA MED DATABASE SEEDING PROCESS ===")
    if not os.path.exists(EXCEL_PATH):
        print(f"Error: Excel file not found at {EXCEL_PATH}")
        return
        
    print(f"Reading Excel file: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    
    # 1. Initialize schema
    init_db()
    
    # 2. Seed Users
    seed_users()
    
    # 3. Seed Products (seeds art_pf first so records can fetch weight weights!)
    seed_products(wb)
    
    # 4. Seed Logs
    seed_production_records(wb)
    
    wb.close()
    print("\nDatabase seeding completed successfully! Ready for desktop app use.")

if __name__ == "__main__":
    main()
